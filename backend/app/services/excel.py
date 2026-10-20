"""Servicio de generación y exportación de presupuestos a hojas de cálculo Excel (.xlsx).

Utiliza la librería oficial openpyxl para crear libros de cálculo estructurados con
estilos corporativos de BIMetrics, formato de moneda COP, fórmulas dinámicas y autoajuste.
"""
import io
from typing import Any, Dict
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def exportar_presupuesto_excel(presupuesto: Dict[str, Any]) -> io.BytesIO:
    """Genera un archivo Excel (.xlsx) en memoria con el detalle del presupuesto.
    
    Args:
        presupuesto: Diccionario con la estructura de presupuesto_completo:
            - nombre_obra: str
            - ciudad: str
            - fecha_calculo: str (ISO)
            - costo_total: float
            - detalle: list[dict] (nombre_insumo, unidad_medida, cantidad, costo_unitario_aplicado, subtotal)
            
    Returns:
        io.BytesIO con el contenido binario del libro .xlsx.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Presupuesto de Obra"
    ws.views.sheetView[0].showGridLines = True

    # -------------------------------------------------------------------------
    # 1. Paleta de Estilos Corporativos BIMetrics
    # -------------------------------------------------------------------------
    font_titulo = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_subtitulo = Font(name="Calibri", size=10, italic=True, color="64748B")
    font_meta_label = Font(name="Calibri", size=10, bold=True, color="1E293B")
    font_meta_val = Font(name="Calibri", size=10, color="334155")
    
    font_encabezado = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_datos = Font(name="Calibri", size=10, color="0F172A")
    font_total = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_total_val = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    font_pie = Font(name="Calibri", size=8.5, italic=True, color="94A3B8")

    fill_banner = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_encabezado = PatternFill(start_color="2456A6", end_color="2456A6", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_total_label = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_total_val = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")

    borde_fino = Side(border_style="thin", color="CBD5E1")
    borde_doble = Side(border_style="double", color="1E293B")
    borde_celda = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)
    borde_total = Border(top=borde_fino, bottom=borde_doble, left=borde_fino, right=borde_fino)

    align_izq = Alignment(horizontal="left", vertical="center")
    align_centro = Alignment(horizontal="center", vertical="center")
    align_der = Alignment(horizontal="right", vertical="center")

    # Formatos numéricos
    fmt_moneda = '"$"#,##0'
    fmt_cantidad = '#,##0.00'

    # -------------------------------------------------------------------------
    # 2. Encabezado y Metadatos del Proyecto
    # -------------------------------------------------------------------------
    # Banner Principal
    ws.merge_cells("A1:F1")
    banner_celda = ws["A1"]
    banner_celda.value = "BIMETRICS  ·  PRESUPUESTO DE OBRA Y ESTIMACIÓN 5D"
    banner_celda.font = font_titulo
    banner_celda.fill = fill_banner
    banner_celda.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Metadatos del Presupuesto
    nombre_obra = presupuesto.get("nombre_obra") or "Proyecto General"
    ciudad = presupuesto.get("ciudad") or "Colombia"
    fecha = str(presupuesto.get("fecha_calculo") or "")[:10]

    ws["A3"] = "Proyecto:"
    ws["A3"].font = font_meta_label
    ws["B3"] = nombre_obra
    ws["B3"].font = font_meta_val

    ws["E3"] = "Ciudad:"
    ws["E3"].font = font_meta_label
    ws["E3"].alignment = align_der
    ws["F3"] = ciudad
    ws["F3"].font = font_meta_val

    ws["A4"] = "Fecha de Cálculo:"
    ws["A4"].font = font_meta_label
    ws["B4"] = fecha
    ws["B4"].font = font_meta_val

    ws["E4"] = "Moneda:"
    ws["E4"].font = font_meta_label
    ws["E4"].alignment = align_der
    ws["F4"] = "COP (Pesos Colombianos)"
    ws["F4"].font = font_meta_val

    # -------------------------------------------------------------------------
    # 3. Tabla de Partidas Presupuestales
    # -------------------------------------------------------------------------
    fila_inicio_tabla = 6
    encabezados = ["#", "DESCRIPCIÓN DEL INSUMO", "UNIDAD", "CANTIDAD", "PRECIO UNITARIO", "SUBTOTAL"]

    for col_idx, enc in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_inicio_tabla, column=col_idx, value=enc)
        celda.font = font_encabezado
        celda.fill = fill_encabezado
        celda.alignment = align_centro if col_idx in (1, 3) else (align_der if col_idx in (4, 5, 6) else align_izq)
        celda.border = borde_celda
    ws.row_dimensions[fila_inicio_tabla].height = 24

    detalles = presupuesto.get("detalle") or []
    fila_actual = fila_inicio_tabla + 1

    if not detalles:
        ws.merge_cells(f"A{fila_actual}:F{fila_actual}")
        celda_vacia = ws[f"A{fila_actual}"]
        celda_vacia.value = "Este presupuesto no contiene partidas registradas."
        celda_vacia.font = font_subtitulo
        celda_vacia.alignment = align_centro
        fila_actual += 1
    else:
        for idx, linea in enumerate(detalles, start=1):
            cantidad = float(linea.get("cantidad") or 0)
            precio_unitario = float(linea.get("costo_unitario_aplicado") or 0)
            
            c_item = ws.cell(row=fila_actual, column=1, value=idx)
            c_nombre = ws.cell(row=fila_actual, column=2, value=linea.get("nombre_insumo"))
            c_unidad = ws.cell(row=fila_actual, column=3, value=linea.get("unidad_medida"))
            c_cant = ws.cell(row=fila_actual, column=4, value=cantidad)
            c_precio = ws.cell(row=fila_actual, column=5, value=precio_unitario)
            c_subtotal = ws.cell(row=fila_actual, column=6, value=f"=D{fila_actual}*E{fila_actual}")

            # Alineaciones y formatos
            c_item.alignment = align_centro
            c_nombre.alignment = align_izq
            c_unidad.alignment = align_centro
            c_cant.alignment = align_der
            c_cant.number_format = fmt_cantidad
            c_precio.alignment = align_der
            c_precio.number_format = fmt_moneda
            c_subtotal.alignment = align_der
            c_subtotal.number_format = fmt_moneda

            # Estilos de celda
            es_par = (idx % 2 == 0)
            for col_i in range(1, 7):
                c = ws.cell(row=fila_actual, column=col_i)
                c.font = font_datos
                c.border = borde_celda
                if es_par:
                    c.fill = fill_zebra

            ws.row_dimensions[fila_actual].height = 20
            fila_actual += 1

    # -------------------------------------------------------------------------
    # 4. Fila de Total Presupuestado
    # -------------------------------------------------------------------------
    ws.merge_cells(f"A{fila_actual}:E{fila_actual}")
    c_tot_label = ws[f"A{fila_actual}"]
    c_tot_label.value = "TOTAL PRESUPUESTADO (COP):"
    c_tot_label.font = font_total
    c_tot_label.alignment = align_der
    c_tot_label.fill = fill_total_label
    c_tot_label.border = borde_total

    # Borde para celdas fusionadas intermedias
    for col_i in range(2, 6):
        c = ws.cell(row=fila_actual, column=col_i)
        c.border = borde_total
        c.fill = fill_total_label

    c_tot_val = ws.cell(row=fila_actual, column=6)
    if detalles:
        c_tot_val.value = f"=SUM(F{fila_inicio_tabla + 1}:F{fila_actual - 1})"
    else:
        c_tot_val.value = float(presupuesto.get("costo_total") or 0)
    c_tot_val.font = font_total_val
    c_tot_val.alignment = align_der
    c_tot_val.number_format = fmt_moneda
    c_tot_val.fill = fill_total_val
    c_tot_val.border = borde_total
    ws.row_dimensions[fila_actual].height = 26

    # -------------------------------------------------------------------------
    # 5. Pie de Página Institucional
    # -------------------------------------------------------------------------
    fila_pie = fila_actual + 2
    ws.merge_cells(f"A{fila_pie}:F{fila_pie}")
    c_pie = ws[f"A{fila_pie}"]
    c_pie.value = "Documento generado por BIMetrics SaaS · Plataforma de Control Presupuestal y Estimación BIM 5D"
    c_pie.font = font_pie
    c_pie.alignment = align_centro

    # -------------------------------------------------------------------------
    # 6. Autoajuste de Ancho de Columnas
    # -------------------------------------------------------------------------
    anchos_minimos = {1: 8, 2: 36, 3: 14, 4: 16, 5: 22, 6: 22}
    for col_idx in range(1, 7):
        col_letra = get_column_letter(col_idx)
        max_len = anchos_minimos.get(col_idx, 12)
        for row_i in range(fila_inicio_tabla, fila_actual + 1):
            val = ws.cell(row=row_i, column=col_idx).value
            if val and not str(val).startswith("="):
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letra].width = min(max_len + 3, 50)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
